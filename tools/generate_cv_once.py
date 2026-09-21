from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import sys

ACCENT='1F4E79'; DARK='1B1F23'; MUTED='5A6570'

RU = {
'name':'НИКОЛАЙ БУКИН','role':'Системный аналитик','city':'Москва','summary_head':'ПРОФЕССИОНАЛЬНЫЙ ПРОФИЛЬ','skills_head':'КЛЮЧЕВЫЕ НАВЫКИ','exp_head':'ОПЫТ РАБОТЫ','edu_head':'ОБРАЗОВАНИЕ И ПРОЕКТЫ',
'summary':'Системный аналитик с практическим опытом работы с требованиями, интеграциями и эксплуатацией корпоративных систем. Совмещаю системный анализ с production-практикой: формализую требования, моделирую процессы, разбираю интеграции и инциденты, тестирую API и готовлю воспроизводимый контекст для разработки. Рабочий стек: REST API, Postman, Python, pytest, FastAPI, Docker, Git, Linux; SQL и реляционные СУБД использую для проверки и анализа данных. Английский - B2.',
'skills':[
('Системный анализ','требования, AS-IS / TO-BE, BPMN, UML, PlantUML, системные сценарии, техническая документация'),
('Данные и интеграции','SQL, PostgreSQL, MariaDB, Microsoft SQL Server, SQLite, REST API, Postman, структуры и маппинг данных'),
('Тестирование и эксплуатация','Python, pytest, FastAPI, Docker, Git, Linux, анализ логов, диагностика инцидентов, 3-я линия поддержки')],
'jobs':[
('АО «ВИМ Инвестиции»','07.2026 - н.в.','Стажер-аналитик, поддержка бизнес-приложений','финансовый сектор, 3-я линия',[
'Автоматизировал повторные регрессионные API-проверки на Python/pytest для сервисов на FastAPI. Эффект: меньше ручного ретестирования, быстрее повторная проверка изменений и ниже риск пропуска регрессий.',
'Сократил цикл первичной диагностики инцидентов: воспроизвожу сценарии, сопоставляю API, логи и данные и локализую слой сбоя до передачи разработчикам. Эффект: меньше уточняющих итераций и быстрее переход к исправлению.',
'Стандартизировал технический контекст по дефектам - шаги, входные данные, ожидаемый/фактический результат и результаты проверки. Эффект: разработчики быстрее воспроизводят проблему, а похожие обращения обрабатываются повторно без диагностики с нуля.'
]),
('Т1','02.2026 - 05.2026','Системный аналитик (стажер)','системная интеграция и разработка ПО',[
'Формализовал функциональные и технические требования и моделировал AS-IS / TO-BE в BPMN. Эффект: более однозначные постановки и меньше риска возврата задач на дополнительные уточнения.',
'Проектировал структуры данных и маппинг, проверял данные SQL-запросами и исследовал REST API в Postman. Эффект: несогласованности между требованиями, данными и интеграциями выявлялись до передачи решения дальше по процессу.'
]),
('РТК ЦОД','03.2025 - 02.2026','Дежурный инженер','центр обработки данных',[
'Проводил мониторинг и первичную диагностику серверной и сетевой инфраструктуры. Эффект: быстрее определялось направление эскалации и снижался риск затяжного простоя. Координировал устранение сбоев, фиксировал состояние систем между сменами и участвовал в профилактическом обслуживании. Эффект: непрерывность работы смен и снижение риска повторных отказов.'
])],
'edu1_b':'Финансовый университет при Правительстве РФ','edu1':' - Прикладная математика и информатика, 2026',
'edu2_b':'ВКР: ','edu2':'Attention Meta-Learner for MOEX Si Futures - 219 529 минутных OHLCV-баров, 29 признаков, несколько ML-моделей, временная валидация и Streamlit.',
'edu3_b':'Дополнительно: ','edu3':'Яндекс Практикум - фронтенд-разработка (2025); Cambridge B2 First (FCE), 2021.',
'about_b':'О себе: ','about':'много лет занимаюсь горными лыжами на высоком уровне, имею серьезный опыт в плавании. Спорт сформировал дисциплину, концентрацию и привычку системно работать на результат.',
'lang':'ru'
}

EN = {
'name':'NIKOLAY BUKIN','role':'Systems Analyst','city':'Moscow','summary_head':'PROFESSIONAL SUMMARY','skills_head':'KEY SKILLS','exp_head':'EXPERIENCE','edu_head':'EDUCATION & PROJECTS',
'summary':'Systems Analyst with hands-on experience in requirements, integrations, and production support of enterprise systems. I combine systems analysis with production operations: formalize requirements, model processes, analyze integrations and incidents, test APIs, and prepare reproducible technical context for development teams. Core stack: REST APIs, Postman, Python, pytest, FastAPI, Docker, Git, Linux; I use SQL and relational databases for data validation and analysis. English - B2.',
'skills':[
('Systems Analysis','requirements, AS-IS / TO-BE, BPMN, UML, PlantUML, system scenarios, technical documentation'),
('Data & Integrations','SQL, PostgreSQL, MariaDB, Microsoft SQL Server, SQLite, REST APIs, Postman, data structures and mapping'),
('Testing & Operations','Python, pytest, FastAPI, Docker, Git, Linux, log analysis, incident diagnostics, L3 support')],
'jobs':[
('VIM Investments JSC','07.2026 - Present','Analyst Intern, Business Application Support','financial sector, L3 support',[
'Automated recurring regression API checks using Python/pytest for FastAPI services. Impact: reduced manual retesting, accelerated change verification, and lowered the risk of missed regressions.',
'Shortened the initial incident diagnosis cycle by reproducing scenarios, correlating APIs, logs, and data, and identifying the failing layer before handoff to developers. Impact: fewer clarification cycles and faster transition to remediation.',
'Standardized technical context for defects - reproduction steps, input data, expected/actual results, and validation findings. Impact: developers reproduce issues faster, while similar incidents can be handled without restarting analysis from scratch.'
]),
('T1','02.2026 - 05.2026','Systems Analyst Intern','systems integration and software development',[
'Formalized functional and technical requirements and modeled AS-IS / TO-BE processes in BPMN. Impact: clearer specifications and lower risk of tasks being returned for additional clarification.',
'Designed data structures and mappings, validated data with SQL queries, and analyzed REST APIs in Postman. Impact: inconsistencies between requirements, data, and integrations were identified before moving solutions further in the delivery process.'
]),
('RTK Data Center','03.2025 - 02.2026','Duty Engineer','data center',[
'Monitored and performed initial diagnostics of server and network infrastructure, coordinated incident resolution, documented system state across shifts, and participated in preventive maintenance. Impact: faster escalation, lower risk of prolonged downtime, and fewer recurring failures.'
])],
'edu1_b':'Financial University under the Government of the Russian Federation','edu1':' - Applied Mathematics and Computer Science, 2026',
'edu2_b':'Thesis: ','edu2':'Attention Meta-Learner for MOEX Si Futures - 219,529 minute OHLCV bars, 29 features, multiple ML models, time-based validation, and Streamlit.',
'edu3_b':'Additional: ','edu3':'Yandex Practicum - Frontend Development (2025); Cambridge B2 First (FCE), 2021.',
'about_b':'About: ','about':'many years of advanced alpine skiing and extensive swimming experience. Sport has developed discipline, focus, and a consistent results-oriented mindset.',
'lang':'en'
}

def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def add_hyperlink(paragraph,text,url):
    rid=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid)
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),ACCENT); rp.append(c)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'none'); rp.append(u)
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'16'); rp.append(sz)
    r.append(rp); t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); paragraph._p.append(h)

def add_floating_picture(paragraph, image_path):
    run=paragraph.add_run(); inline_shape=run.add_picture(str(image_path), width=Cm(2.7))
    inline=inline_shape._inline
    inline.tag=qn('wp:anchor')
    for k,v in {'distT':'0','distB':'0','distL':'90000','distR':'90000','simplePos':'0','relativeHeight':'251659264','behindDoc':'0','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items(): inline.set(k,v)
    for k in list(inline.attrib.keys()):
        if k not in {'distT','distB','distL','distR','simplePos','relativeHeight','behindDoc','locked','layoutInCell','allowOverlap'}:
            del inline.attrib[k]
    simple=OxmlElement('wp:simplePos'); simple.set('x','0'); simple.set('y','0')
    ph=OxmlElement('wp:positionH'); ph.set('relativeFrom','page'); pho=OxmlElement('wp:posOffset'); pho.text='6210000'; ph.append(pho)
    pv=OxmlElement('wp:positionV'); pv.set('relativeFrom','page'); pvo=OxmlElement('wp:posOffset'); pvo.text='198000'; pv.append(pvo)
    wrap=OxmlElement('wp:wrapSquare'); wrap.set('wrapText','bothSides')
    inline.insert(0,simple); inline.insert(1,ph); inline.insert(2,pv)
    idx=0
    for i,ch in enumerate(inline):
        if ch.tag==qn('wp:effectExtent'): idx=i+1
    inline.insert(idx,wrap)

def para_fmt(p,before=0,after=0,keep=False,keep_next=False):
    pf=p.paragraph_format; pf.line_spacing=1.0583333333333333; pf.space_before=Pt(before); pf.space_after=Pt(after); pf.keep_together=keep; pf.keep_with_next=keep_next

def section_title(doc,text):
    p=doc.add_paragraph(); para_fmt(p,2,0.8,True,True)
    r=p.add_run(text); r.bold=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(ACCENT)
    pPr=p._p.get_or_add_pPr(); bdr=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom')
    bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'6'); bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),ACCENT); bdr.append(bottom); pPr.append(bdr)

def bullet(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(0.28); p.paragraph_format.first_line_indent=Cm(-0.18); para_fmt(p,0,0.15,True,False)
    r=p.add_run('• '); r.bold=True; r.font.color.rgb=RGBColor.from_string(ACCENT); p.add_run(text)

def make_doc(data, out_path, photo_path):
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21.0); sec.page_height=Cm(29.7); sec.top_margin=Cm(0.75); sec.bottom_margin=Cm(0.72); sec.left_margin=Cm(1.05); sec.right_margin=Cm(1.05); sec.header_distance=Cm(0.25); sec.footer_distance=Cm(0.25)
    st=doc.styles['Normal']; st.font.name='Arial'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); st.font.size=Pt(8.0); st.paragraph_format.line_spacing=1.0583333333333333; st.paragraph_format.space_after=Pt(0)
    p=doc.add_paragraph(); para_fmt(p,0,0.2); r=p.add_run(data['name']); r.bold=True; r.font.size=Pt(18.5); r.font.color.rgb=RGBColor.from_string(DARK); add_floating_picture(p,photo_path)
    p=doc.add_paragraph(); para_fmt(p,0,0.35); r=p.add_run(data['role']); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor.from_string(ACCENT)
    p=doc.add_paragraph(); para_fmt(p); r=p.add_run(f"{data['city']}  |  +7 (996) 864-69-42  |  bkolyabu@gmail.com"); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)
    p=doc.add_paragraph(); para_fmt(p,0,1.2); r=p.add_run('Telegram: @Boo4kin  |  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED); add_hyperlink(p,'GitHub','https://github.com/Boo4kin')
    section_title(doc,data['summary_head'])
    p=doc.add_paragraph(); para_fmt(p,0,0.5,True,False); p.add_run(data['summary'])
    section_title(doc,data['skills_head'])
    for lab,txt in data['skills']:
        p=doc.add_paragraph(); para_fmt(p,0,0.05,True,False); r=p.add_run(lab+': '); r.bold=True; r.font.color.rgb=RGBColor.from_string(DARK); p.add_run(txt)
    section_title(doc,data['exp_head'])
    for company,dates,role,context,bullets in data['jobs']:
        t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False
        for c in t.rows[0].cells: c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; set_cell_margins(c,0,0,0,0)
        p=t.cell(0,0).paragraphs[0]; para_fmt(p); r=p.add_run(company); r.bold=True; r.font.size=Pt(9.2); r.font.color.rgb=RGBColor.from_string(DARK)
        p=t.cell(0,1).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; para_fmt(p); r=p.add_run(dates); r.bold=True; r.font.size=Pt(7.8); r.font.color.rgb=RGBColor.from_string(MUTED)
        p=doc.add_paragraph(); para_fmt(p,0,0.05,False,True); r=p.add_run(role); r.bold=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(ACCENT); r=p.add_run('  |  '+context); r.italic=True; r.font.size=Pt(7.5); r.font.color.rgb=RGBColor.from_string(MUTED)
        for b in bullets: bullet(doc,b)
    section_title(doc,data['edu_head'])
    for b,t in [(data['edu1_b'],data['edu1']),(data['edu2_b'],data['edu2']),(data['edu3_b'],data['edu3'])]:
        p=doc.add_paragraph(); para_fmt(p,0,0.05,True,False); r=p.add_run(b); r.bold=True; p.add_run(t)
    p=doc.add_paragraph(); para_fmt(p,0.3,0,True,False); r=p.add_run(data['about_b']); r.bold=True; p.add_run(data['about'])
    sec.footer.paragraphs[0].text=''
    props=doc.core_properties; props.title=f"Nikolay Bukin - {'Systems Analyst' if data['lang']=='en' else 'Системный аналитик'}"; props.subject='Resume' if data['lang']=='en' else 'Резюме'; props.author='Nikolay Bukin' if data['lang']=='en' else 'Николай Букин'
    doc.save(out_path)

if __name__=='__main__':
    out=Path(sys.argv[1]) if len(sys.argv)>1 else Path('cvbuild'); out.mkdir(parents=True,exist_ok=True)
    photo=Path(sys.argv[2]) if len(sys.argv)>2 else Path('assets/img/nikolay-bukin-original.jpg')
    make_doc(RU,out/'Nikolay-Bukin-CV-RU.docx',photo)
    make_doc(EN,out/'Nikolay-Bukin-CV-EN.docx',photo)
